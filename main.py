#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Creates a wave specified by the timetable in the excel file.

Raises:
    ValueError: The excel table must contain information for all wave chambers. The number of chambers N_CHAMBERS is specified in setup.py.
    ValueError: The excel table must describe a waveset composed of (possibly multiple) waves, where each wave is fully characterized by four phases. The time intervals for all four phases for each wave in the waveset have to be provided.
"""


from    openpyxl    import load_workbook
import  pandas      as pd
import  revpimodio2
from    setup       import N_CHAMBERS, PATH, TABLENAME
import  waves

__author__      = "Maximilian Detering"
__copyright__   = "Copyright (c) 2023 Maximilian Detering"
__version__     = "1.0"


def main():
    # load table with name TABLE (see setup.py) from excel file located at PATH (see setup.py)
    wb = load_workbook(filename = PATH)
    sheet = wb['Sheet1']
    lookup_table = sheet.tables[TABLENAME]

    # Access the data in the table range
    data = sheet[lookup_table.ref]
    table_values = []
    column_labels = []
    row_labels = []

    # Loop through each row and get the values in the cells
    for i, row in enumerate(data):
        # get column labels
        if i == 0:
            for label in row[1:]:
                column_labels.append(label.value)
        # get row labels and table data
        else:
            # get row labels
            row_labels.append(row[0].value)
            # get table data
            cells = []
            for cell in row[1:]:
                cells.append(cell.value)
            table_values.append(cells)

    # create a pandas dataframe from the rows_list.
    df = pd.DataFrame(data=table_values, index=row_labels, columns=column_labels, dtype=float)

    # check if file contains information for all wave chambers
    if len(df.index) != N_CHAMBERS:
         raise ValueError("Table does not contain information for all chambers or for too many chambers!")

    # get number of wave sets specified in file and perform sanity check
    if len(df.columns) % 4 != 0:
           raise ValueError("Table contains incomplete wavesets! Length of columns must be multiple of 4.")

    # initialize the RevPiModIO object to later control the IO pins
    rpi = revpimodio2.RevPiModIO(autorefresh=True)

    # generate waves according to the timetable specified in the excel file
    waves.create_waveset(rpi, df.values)

    # close all valves again
    waves.clean_up(rpi)



if __name__ == "__main__":
    main()